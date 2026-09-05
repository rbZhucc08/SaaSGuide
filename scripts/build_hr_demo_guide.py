from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "SaaSGuide_V2_HR_演示引导.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), color)
        borders.append(tag)


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    header = table.rows[0]
    for index, text in enumerate(headers):
        cell = header.cells[index]
        cell.text = text
        set_cell_shading(cell, "27324A")
        set_cell_border(cell)
        set_cell_margin(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(9.5)
        if widths:
            cell.width = Cm(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, text in enumerate(values):
            cells[index].text = str(text)
            set_cell_border(cells[index])
            set_cell_margin(cells[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cells[index], "F3F6FA")
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.1
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
            if widths:
                cells[index].width = Cm(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_steps(doc, items):
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.8)
        paragraph.paragraph_format.first_line_indent = Cm(-0.8)
        paragraph.add_run(f"{index}.  {item}")


def new_page(doc, title, intro=None):
    doc.add_page_break()
    heading = doc.add_heading(title, level=1)
    keep_with_next(heading)
    if intro:
        doc.add_paragraph(intro)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.35
    for name, size in (("Title", 30), ("Heading 1", 20), ("Heading 2", 13)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(18)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    styles["Heading 1"].paragraph_format.space_before = Pt(0)
    styles["Heading 1"].paragraph_format.space_after = Pt(12)
    styles["Heading 2"].paragraph_format.space_before = Pt(12)
    styles["Heading 2"].paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph(style="Title")
    title.add_run("SaaSGuide V2 HR 演示引导")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = subtitle.add_run("个人学习 Demo 的 5 至 8 分钟面试讲解稿")
    run.font.size = Pt(16)
    run.font.bold = True
    doc.add_paragraph("这份文档用于在面试或作品集沟通中演示 SaaSGuide V2。演示的重点不是展示一组漂亮的固定页面，而是说明输入如何被校验、证据如何保留、哪些步骤必须人工确认，以及哪些能力仍未验证。")
    doc.add_heading("30 秒项目定位", level=1)
    doc.add_paragraph("SaaSGuide V2 是一个本地运行的项目风险学习 Demo。它把项目任务 XLSX、文本证据、自建知识文档、风险规则、人工决策、SQLite 行动记录和报告导出串成一个可追溯流程。我用确定性代码处理可计算和可校验的部分，把不确定判断留给人工确认，并明确区分已实现能力、固定模拟评测和真实企业效果。")
    doc.add_heading("开场真实性声明", level=2)
    doc.add_paragraph("这是个人学习 Demo，使用虚构 SaaS、自建规则、模拟样本或用户主动导入的本地文件。没有真实客户、真实企业系统写回、生产部署或业务收益数据。固定测试通过只证明给定条件下的功能，不证明真实企业准确率。")
    add_table(doc, ["本次能证明", "本次不能证明"], [["本地端到端工作流、结构校验、来源追踪、人工确认和测试习惯", "真实客户采用、风险降低、生产稳定性或商业价值"], ["普通 XLSX、TXT、Markdown、DOCX、普通 PDF 的已验收边界", "真实 OCR、语音识别、向量 RAG、外部连接器和多人权限"]], [8.0, 8.0])

    new_page(doc, "演示总路线", "建议控制在 5 至 8 分钟。若面试官只给 3 分钟，优先演示数据源、风险雷达和人工确认。")
    add_table(doc, ["时间", "页面", "要展示的动作", "要说明的结论"], [
        ["0:00–0:40", "工作台 /", "指出指标来自本地已保存状态", "首页不再使用固定五条风险填充"],
        ["0:40–2:10", "数据源 /data-sources", "加载有效样本并核对开放式列映射", "建议不是限制；服务器会重新校验"],
        ["2:10–3:30", "风险雷达 /risk-radar", "运行规则并展开来源证据", "候选风险不是正式结论"],
        ["3:30–4:30", "文本证据 /evidence-intake", "展示来源位置和人工核对", "文件内容按不可信输入处理"],
        ["4:30–5:30", "知识库 /knowledge-base", "提问、查看引用或拒答", "有依据才答，版本状态可追溯"],
        ["5:30–6:30", "行动跟踪 /action-tracker", "展示人工确认和状态事件", "规则和模型不自动指派真实人员"],
        ["6:30–7:30", "报告 /reports", "核对指标并导出 CSV 或 XLSX", "数字由 Python 计算，不由模型心算"],
        ["7:30–8:00", "回到工作台", "总结输入到报告的闭环", "再次说明未验证边界"],
    ], [2.0, 3.2, 5.2, 5.6])
    doc.add_heading("演示前准备", level=2)
    add_bullets(doc, ["运行 start_v2.ps1，并确认 http://127.0.0.1:4173/ 可打开。", "准备 data/samples/valid_project_tasks_cn.xlsx；真实文件选择若不稳定，可使用页面明确标注的内置模拟样本。", "不要预先承诺模型、OCR、语音或外部系统会现场成功；这些不在当前已验证范围。"])

    new_page(doc, "第一段 数据接入与映射", "这一段用于证明系统不是只展示固定结果，而是先处理来源、字段和错误。")
    doc.add_heading("点击顺序", level=2)
    add_steps(doc, ["从工作台点击导入项目数据。", "选择 valid_project_tasks_cn.xlsx，或选择内置有效样本后点击加载。", "指出左侧是统一字段，右侧是可自由输入的原始列名；点入输入框可看到原表头建议。", "把一个必填字段改为不存在的列名，点击按当前映射重新校验，展示 mapping_unknown_column。", "恢复正确列名并重新校验；确认只有校验通过后才能导入。", "若尚未导入过相同内容，点击确认导入并展示来源哈希、统一 JSON 与任务数量。"])
    doc.add_heading("讲解重点", level=2)
    add_bullets(doc, ["Auto suggestion 自动建议只减少输入成本，不代表系统知道企业字段含义。", "Validation 校验由服务端再次执行，不能只相信浏览器输入。", "Provenance 来源追踪保留文件名、SHA-256 和原表行号，方便回到证据。", "Human confirmation 人工确认是写入正式本地数据前的闸门。"])
    doc.add_heading("不要说", level=2)
    doc.add_paragraph("不要说系统兼容所有 Excel，也不要把内置样本的成功说成真实企业文件验证。当前只对已测试的 XLSX 范围负责。")

    new_page(doc, "第二段 风险候选与证据", "这一段展示规则如何从统一任务数据生成可解释的 Risk Candidate 风险候选，而不是自动生成业务事实。")
    doc.add_heading("风险雷达", level=2)
    add_steps(doc, ["运行最近导入数据扫描或固定样本扫描。", "展开一个候选，指出规则编号、源任务、源表行号和字段值。", "解释按项目、任务和风险类型去重。", "选择确认、观察、驳回或误报，并说明它只追加人工审计记录。"])
    add_table(doc, ["概念", "在项目中的含义", "容易混淆的错误说法"], [
        ["Risk Signal", "截止日期、状态、完成度或依赖触发的可计算信号", "AI 已经发现真实风险"],
        ["Risk Candidate", "带证据、等待人工判断的候选", "候选就是正式风险"],
        ["Ground Truth", "固定评测集里人工预先标注的答案", "代表行业真实准确率"],
        ["Human Decision", "确认、观察、驳回或误报的追加审计", "系统自动替人承担责任"],
    ], [3.2, 7.1, 5.3])
    doc.add_heading("固定评测的正确表述", level=2)
    doc.add_paragraph("P2 固定模拟评测集的 Precision 和 Recall 均为 83.33%。这说明规则在这组自建样本上有可见的误报和漏报，不是企业部署效果，也不能外推到其他项目。")

    new_page(doc, "第三段 文本 知识 行动与报告", "这四个页面共同说明证据、规则、人工责任和可交付输出如何衔接。")
    add_table(doc, ["页面", "演示动作", "你要说明的边界"], [
        ["文本证据", "解析 TXT、Markdown 或 DOCX，查看段落或表格位置，保存人工核对", "真实模型抽取未验证；提示词注入文字不会变成系统指令"],
        ["知识库", "提出有依据问题并查看版本引用，再问一个资料外问题观察拒答", "这是确定性检索底座，不是向量或模型 RAG"],
        ["行动跟踪", "展示人工确认创建、受限状态转换和追加事件", "SQLite 是单机 Demo，不是企业权限或并发系统"],
        ["报告", "核对指标口径，下载 CSV 或三表 XLSX", "数字来自 Python 确定性计算；固定数据不等于业务成果"],
        ["输入实验室", "展示普通 PDF、OCR_REQUIRED、WAV 元数据和 simulation 适配器", "未调用真实 OCR、语音识别或外部连接器"],
    ], [3.0, 6.4, 6.2])
    doc.add_heading("闭环总结话术", level=2)
    doc.add_paragraph("这个项目的主线不是让模型自由决定，而是把文件接入、确定性规则、知识引用、人工决策、行动审计和报告拆成可验证步骤。每一步都有明确输入输出；不确定信息停在待确认状态，不能静默写成事实。")

    new_page(doc, "技术串联讲法", "面试官追问实现方式时，用下面的层次回答，避免把 Skill、Agent、Workflow 和 Tool 混为一谈。")
    add_table(doc, ["层次", "通俗解释", "本项目例子"], [
        ["Skill", "完成某类任务的操作规范和质量门槛", "文档生成要求 DOCX 渲染后逐页检查"],
        ["Agent", "理解目标、选择步骤并协调工具的执行者", "当前开发由一个主 Codex agent 执行；产品架构也只设一个 Orchestrator 概念"],
        ["Workflow", "业务状态和工具调用的固定顺序", "上传→映射→校验→预览→人工确认→保存"],
        ["Tool", "输入输出明确的函数或服务", "parse_xlsx、规则扫描、知识检索、指标计算"],
        ["Test", "验证代码在给定条件下是否符合预期", "单元测试、接口测试、浏览器交互和文档视觉检查"],
    ], [2.6, 6.0, 7.0])
    doc.add_heading("一句话串联", level=2)
    doc.add_paragraph("用户操作进入 Workflow，主流程调用确定性 Tool；遇到需要判断的内容时保留证据并进入人工确认；结果再写入 SQLite 或文件审计，最后由报告工具汇总。Skill 约束的是开发和产物质量，不是产品页面里的一个机器人。")
    doc.add_heading("当前多 Agent 事实", level=2)
    doc.add_paragraph("本次产品化修订没有创建子 agent，也没有多个 agent 自由讨论。把产品中的 Orchestrator Agent 画成架构概念，不应表述为已经实现了多智能体系统。")

    new_page(doc, "常见 HR 追问", "回答时先给事实，再给限制，最后说明下一步。")
    add_table(doc, ["问题", "简短回答"], [
        ["这是正式产品吗", "不是。它是本地个人学习 Demo，用来证明我能把需求拆成可运行、可测试、可解释的流程。"],
        ["AI 在哪里", "V1 有真实 DeepSeek ASK/PLAN 调用证据；V2 的文本抽取和知识检索当前以确定性底座为主，我没有把未调用的模型能力写成完成。"],
        ["为什么不用模型直接判断风险", "日期、状态、依赖等先用规则更可复现。模型适合提问、解释和草拟，但正式状态必须由人确认。"],
        ["83.33% 准确吗", "它是自建固定模拟集上的 Precision 和 Recall，不是现实准确率；我保留了误报和漏报，避免把测试包装成业务效果。"],
        ["为什么用 SQLite", "它足够验证本地关联、状态机和事件审计，但不代表多租户、权限、并发或灾备。"],
        ["你亲自做了什么", "我负责需求取舍、验收边界和结果判断，并能按数据接入、规则、人工确认、审计和报告解释代码关系；开发过程使用 Codex 辅助实现和测试。"],
        ["下一步做什么", "先用新的非同源样本做盲测，再选择一个真实模型抽取或外部系统只读连接进行独立验收，而不是同时堆功能。"],
    ], [5.0, 10.6])

    new_page(doc, "结束检查清单", "结束前用 20 秒确认自己没有把计划、模拟或局部验证说成事实。")
    add_bullets(doc, [
        "我是否明确说了个人学习 Demo、虚构 SaaS 和本地数据边界。",
        "我是否展示了输入、处理、输出、验证和人工确认，而不只是展示页面。",
        "我是否把固定评测指标限定在自建样本上。",
        "我是否避免声称真实 OCR、语音识别、向量 RAG、外部连接器、企业权限和公开部署。",
        "我是否能解释一个真实 Bug，例如下拉框封闭输入、旧验收页混入产品或 SQLite 连接未关闭。",
        "我是否说明 Codex 是开发协作工具，而不是替我拥有真实工作经验。",
    ])
    doc.add_heading("最终结束语", level=2)
    doc.add_paragraph("SaaSGuide V2 当前最有价值的证据，是我能识别错误产品前提并改正：产品页只保留用户任务，演示说明移到独立文档，固定样板退出正式首页，开放输入由服务端重新校验。它展示的是可验证的工程和产品判断，不是虚构的企业成绩。")

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("SaaSGuide V2 HR 演示引导  |  个人学习 Demo")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(110, 118, 136)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
