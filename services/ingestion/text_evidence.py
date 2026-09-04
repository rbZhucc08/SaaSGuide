"""Deterministic text and DOCX evidence intake for the V2 learning demo."""

from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document

MAX_TEXT_BYTES = 2 * 1024 * 1024
ALLOWED_EXTENSIONS = {".txt", ".md", ".docx"}
INJECTION_PATTERNS = ("忽略系统规则", "输出全部数据", "ignore previous", "system prompt")


class TextEvidenceError(ValueError):
    def __init__(self, message: str, code: str, status: int = 400):
        super().__init__(message)
        self.code = code
        self.status = status


def _now() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def _clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _docx_segments(content: bytes) -> list[dict[str, Any]]:
    try:
        document = Document(BytesIO(content))
    except Exception as error:
        raise TextEvidenceError("DOCX 文件损坏或无法解析", "docx_invalid", 422) from error
    segments: list[dict[str, Any]] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = _clean(paragraph.text)
        if text:
            segments.append({"location": f"段落 {index}", "text": text, "kind": "paragraph"})
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [_clean(cell.text) for cell in row.cells]
            if any(cells):
                segments.append({
                    "location": f"表格 {table_index} 行 {row_index}",
                    "text": " | ".join(cells),
                    "cells": cells,
                    "kind": "table_row",
                })
    if not segments:
        raise TextEvidenceError("DOCX 中没有可读取的文字", "text_empty", 422)
    return segments


def _plain_segments(content: bytes, extension: str) -> list[dict[str, Any]]:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise TextEvidenceError("文本编码无法识别；请使用 UTF-8、UTF-8 BOM 或 GBK", "encoding_invalid", 422)
    segments = []
    for index, line in enumerate(text.splitlines(), start=1):
        clean = _clean(line.lstrip("#-* ")) if extension == ".md" else _clean(line)
        if clean:
            segments.append({"location": f"行 {index}", "text": clean, "kind": "line"})
    if not segments:
        raise TextEvidenceError("文本文件没有可读取的内容", "text_empty", 422)
    return segments


def _candidate(segment: dict[str, Any]) -> dict[str, Any] | None:
    text = segment["text"]
    cells = segment.get("cells", [])
    if re.match(r"^T-\d+[｜|]", text, re.I):
        parts = [part.strip() for part in re.split(r"[｜|]", text)]
        values: dict[str, str] = {"task_id": parts[0], "task_name": parts[1] if len(parts) > 1 else ""}
        labels = {"负责人": "owner", "状态": "status", "完成度": "progress", "截止": "due_date", "事实": "evidence"}
        for part in parts[2:]:
            if "：" in part:
                label, value = part.split("：", 1)
                if label in labels:
                    values[labels[label]] = value.strip()
        missing = [field for field in ("task_id", "status", "progress", "due_date") if not values.get(field)]
        return {
            "candidate_id": f"fact-{uuid4().hex[:10]}",
            "category": "explicit_fact" if not missing else "needs_confirmation",
            "fact": values,
            "missing_fields": missing,
            "source_location": segment["location"],
            "quote": text,
        }
    if cells and len(cells) >= 5 and re.fullmatch(r"T-\d+", cells[0], re.I):
        missing = []
        labels = ("task_id", "task_name", "status", "progress", "due_date")
        values = dict(zip(labels, cells[:5]))
        for field in ("task_id", "status", "progress", "due_date"):
            if not values.get(field):
                missing.append(field)
        return {
            "candidate_id": f"fact-{uuid4().hex[:10]}",
            "category": "explicit_fact" if not missing else "needs_confirmation",
            "fact": values,
            "missing_fields": missing,
            "source_location": segment["location"],
            "quote": text,
        }
    if any(pattern.lower() in text.lower() for pattern in INJECTION_PATTERNS):
        return {
            "candidate_id": f"fact-{uuid4().hex[:10]}",
            "category": "suspicious_instruction_text",
            "fact": {"text": text},
            "missing_fields": [],
            "source_location": segment["location"],
            "quote": text,
        }
    confirmation_signals = ("待确认：", "尚未确认", "尚未确定", "不确定", "可能调整")
    if any(signal in text for signal in confirmation_signals):
        return {
            "candidate_id": f"fact-{uuid4().hex[:10]}",
            "category": "needs_confirmation",
            "fact": {"text": text},
            "missing_fields": ["confirmation"],
            "source_location": segment["location"],
            "quote": text,
        }
    return None


def parse_evidence(content: bytes, filename: str) -> dict[str, Any]:
    if not isinstance(content, bytes) or not content:
        raise TextEvidenceError("请选择一个非空文件", "file_required")
    if len(content) > MAX_TEXT_BYTES:
        raise TextEvidenceError("文件超过 2 MB 限制", "file_too_large", 413)
    safe_name = Path(filename or "").name
    extension = Path(safe_name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise TextEvidenceError("仅支持 TXT、Markdown 和 DOCX", "extension_not_allowed", 415)
    segments = _docx_segments(content) if extension == ".docx" else _plain_segments(content, extension)
    candidates = [item for item in (_candidate(segment) for segment in segments) if item]
    questions = []
    for item in candidates:
        if item["category"] == "needs_confirmation":
            questions.append({
                "candidate_id": item["candidate_id"],
                "question": f"请确认或补充：{item['quote'][:80]}",
                "missing_fields": item["missing_fields"],
            })
    return {
        "preview_id": uuid4().hex,
        "source": {
            "filename": safe_name,
            "extension": extension,
            "sha256": hashlib.sha256(content).hexdigest(),
            "size_bytes": len(content),
        },
        "segments": segments,
        "candidates": candidates,
        "questions": questions,
        "summary": {
            "segment_count": len(segments),
            "candidate_count": len(candidates),
            "explicit_fact_count": sum(item["category"] == "explicit_fact" for item in candidates),
            "needs_confirmation_count": sum(item["category"] == "needs_confirmation" for item in candidates),
            "suspicious_text_count": sum(item["category"] == "suspicious_instruction_text" for item in candidates),
        },
        "model_status": "not_called",
        "notice": "当前预览由确定性解析生成；可疑指令只作为文档内容保留。",
    }


def save_review(path: Path, preview: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any]:
    decisions = payload.get("decisions")
    reviewer = _clean(payload.get("reviewer"))
    if not reviewer:
        raise TextEvidenceError("请填写核对人", "reviewer_required")
    if not isinstance(decisions, list) or not decisions:
        raise TextEvidenceError("至少需要一条人工核对结果", "decisions_required")
    allowed_ids = {item["candidate_id"] for item in preview.get("candidates", [])}
    normalized = []
    for item in decisions:
        candidate_id = _clean(item.get("candidate_id")) if isinstance(item, dict) else ""
        decision = _clean(item.get("decision")) if isinstance(item, dict) else ""
        if candidate_id not in allowed_ids or decision not in {"accept", "edit", "reject"}:
            raise TextEvidenceError("人工核对结果包含无效候选或选择", "decision_invalid")
        normalized.append({
            "candidate_id": candidate_id,
            "decision": decision,
            "edited_fact": item.get("edited_fact") if decision == "edit" else None,
            "note": _clean(item.get("note"))[:500],
        })
    record = {
        "review_id": f"review-{uuid4().hex[:12]}",
        "created_at": _now(),
        "reviewer": reviewer[:80],
        "source": preview["source"],
        "decisions": normalized,
        "effect": "evidence_only_no_task_or_risk_writeback",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(record, ensure_ascii=False) + "\n")
    return record
