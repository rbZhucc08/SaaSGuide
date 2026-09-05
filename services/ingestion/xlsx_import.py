"""Parse, validate, preview, and confirm Phase 1 XLSX project-task imports."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook


MAX_XLSX_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_ZIP_ENTRIES = 1000
MAX_ROWS = 1000
MAX_COLUMNS = 100
MAX_PREVIEW_ROWS = 20
SCHEMA_VERSION = "2.0-phase1"

FIELD_DEFINITIONS = [
    {"key": "project_id", "label": "项目编号", "required": True},
    {"key": "project_name", "label": "项目名称", "required": True},
    {"key": "task_id", "label": "任务编号", "required": True},
    {"key": "task_name", "label": "任务名称", "required": True},
    {"key": "owner", "label": "负责人", "required": True},
    {"key": "start_date", "label": "开始日期", "required": True},
    {"key": "due_date", "label": "截止日期", "required": True},
    {"key": "status", "label": "状态", "required": True},
    {"key": "priority", "label": "优先级", "required": False},
    {"key": "dependency_ids", "label": "前置任务", "required": False},
    {"key": "progress_percent", "label": "完成百分比", "required": False},
    {"key": "effort_hours", "label": "工作量（小时）", "required": False},
    {"key": "update_text", "label": "最新进展", "required": False},
    {"key": "updated_at", "label": "更新时间", "required": False},
]

FIELD_ALIASES = {
    "project_id": {"项目编号", "项目id", "projectid", "project_id"},
    "project_name": {"项目名称", "项目名", "projectname", "project_name"},
    "task_id": {"任务编号", "任务id", "taskid", "task_id"},
    "task_name": {"任务名称", "任务名", "taskname", "task_name", "title"},
    "owner": {"负责人", "责任人", "owner", "assignee"},
    "start_date": {"开始日期", "计划开始日期", "startdate", "start_date"},
    "due_date": {"截止日期", "结束日期", "计划结束日期", "duedate", "due_date"},
    "status": {"状态", "任务状态", "status"},
    "priority": {"优先级", "priority"},
    "dependency_ids": {"前置任务", "依赖任务", "dependencies", "dependencyids", "dependency_ids"},
    "progress_percent": {"完成百分比", "完成度", "进度", "progress", "progresspercent", "progress_percent"},
    "effort_hours": {"工作量小时", "工作量（小时）", "工时", "efforthours", "effort_hours"},
    "update_text": {"最新进展", "任务更新", "updatetext", "update_text"},
    "updated_at": {"更新时间", "更新日期", "updatedat", "updated_at"},
}

STATUS_ALIASES = {
    "未开始": "未开始",
    "notstarted": "未开始",
    "todo": "未开始",
    "进行中": "进行中",
    "inprogress": "进行中",
    "阻塞": "阻塞",
    "blocked": "阻塞",
    "已完成": "已完成",
    "completed": "已完成",
    "done": "已完成",
    "已取消": "已取消",
    "cancelled": "已取消",
    "canceled": "已取消",
}


class IngestionError(ValueError):
    def __init__(self, code: str, message: str, status: int = 400) -> None:
        super().__init__(message)
        self.code = code
        self.status = status


@dataclass(frozen=True)
class ParsedWorkbook:
    source_name: str
    sheet_name: str
    headers: list[str]
    rows: list[dict[str, Any]]
    sha256: str


def _safe_filename(filename: str) -> str:
    name = Path(str(filename or "")).name.strip()
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name)
    return name[:120] or "project-tasks.xlsx"


def _header_key(value: str) -> str:
    return re.sub(r"[\s_\-（）()]", "", value.strip().lower())


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _error(code: str, message: str, row: int | None = None, field: str | None = None) -> dict[str, Any]:
    item: dict[str, Any] = {"code": code, "message": message}
    if row is not None:
        item["row"] = row
    if field is not None:
        item["field"] = field
    return item


def inspect_xlsx_archive(path: Path) -> None:
    if not zipfile.is_zipfile(path):
        raise IngestionError("invalid_xlsx", "文件不是有效的 XLSX 工作簿")
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise IngestionError("xlsx_too_complex", "XLSX 内部文件数量超过 1000 个")
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise IngestionError("xlsx_uncompressed_too_large", "XLSX 解压后内容超过 20 MB")
    except zipfile.BadZipFile as error:
        raise IngestionError("invalid_xlsx", "文件不是有效的 XLSX 工作簿") from error


def parse_xlsx(path: Path, source_name: str | None = None) -> ParsedWorkbook:
    if path.suffix.lower() != ".xlsx":
        raise IngestionError("unsupported_file_type", "当前只支持 .xlsx 文件")
    size = path.stat().st_size
    if size == 0:
        raise IngestionError("empty_file", "XLSX 文件不能为空")
    if size > MAX_XLSX_BYTES:
        raise IngestionError("file_too_large", "XLSX 文件不能超过 2 MB", 413)
    inspect_xlsx_archive(path)

    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except Exception as error:
        raise IngestionError("invalid_xlsx", "XLSX 工作簿无法读取") from error

    try:
        visible_sheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        if len(visible_sheets) != 1:
            raise IngestionError("visible_sheet_count", "第一阶段要求且只允许 1 个可见工作表")
        sheet = visible_sheets[0]
        sheet_name = sheet.title
        raw_rows = []
        for row_number, cells in enumerate(sheet.iter_rows(), start=1):
            if row_number > MAX_ROWS + 1:
                raise IngestionError("too_many_rows", "工作表不能超过 1000 条数据行")
            trimmed_cells = list(cells)
            while trimmed_cells and trimmed_cells[-1].value in (None, ""):
                trimmed_cells.pop()
            if len(trimmed_cells) > MAX_COLUMNS:
                raise IngestionError("too_many_columns", "工作表不能超过 100 列")
            raw_rows.append(tuple(trimmed_cells))
        if not raw_rows:
            raise IngestionError("empty_sheet", "工作表没有标题行")

        header_cells = raw_rows[0]
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in header_cells]
        while headers and not headers[-1]:
            headers.pop()
        if not headers:
            raise IngestionError("empty_header", "第一行必须包含列名")
        if any(not header for header in headers):
            raise IngestionError("empty_header", "第一行不能包含空列名")
        normalized_headers = [_header_key(header) for header in headers]
        if len(normalized_headers) != len(set(normalized_headers)):
            raise IngestionError("duplicate_header", "第一行不能包含重复列名")

        rows: list[dict[str, Any]] = []
        for excel_row, cells in enumerate(raw_rows[1:], start=2):
            selected_cells = cells[: len(headers)]
            nonempty = [cell for cell in selected_cells if cell.value not in (None, "")]
            if not nonempty:
                continue
            formula_cell = next((cell for cell in selected_cells if cell.data_type == "f"), None)
            if formula_cell is not None:
                raise IngestionError(
                    "formula_not_supported",
                    f"第 {excel_row} 行包含公式；请先在 Excel 中转换为值",
                )
            rows.append(
                {
                    "source_row": excel_row,
                    "values": {
                        header: _json_value(selected_cells[index].value)
                        for index, header in enumerate(headers)
                    },
                }
            )
        if not rows:
            raise IngestionError("empty_dataset", "工作表没有可导入的数据行")
    finally:
        workbook.close()

    return ParsedWorkbook(
        source_name=_safe_filename(source_name or path.name),
        sheet_name=sheet_name,
        headers=headers,
        rows=rows,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
    )


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    normalized = {_header_key(header): header for header in headers}
    result: dict[str, str] = {}
    for definition in FIELD_DEFINITIONS:
        aliases = {_header_key(alias) for alias in FIELD_ALIASES[definition["key"]]}
        result[definition["key"]] = next(
            (original for key, original in normalized.items() if key in aliases),
            "",
        )
    return result


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _parse_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        pass
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y年%m月%d日"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            pass
    return None


def _parse_datetime(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).isoformat()
    except ValueError:
        parsed_date = _parse_date(text)
        return parsed_date


def _parse_number(value: Any) -> float | None:
    if value in (None, "") or isinstance(value, bool):
        return None
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return None


def _split_dependencies(value: Any) -> list[str]:
    return [part.strip() for part in re.split(r"[,，;；]", _text(value)) if part.strip()]


def validate_and_normalize(parsed: ParsedWorkbook, mapping: dict[str, Any]) -> dict[str, Any]:
    errors: list[dict[str, Any]] = []
    clean_mapping: dict[str, str] = {}
    header_set = set(parsed.headers)
    required_keys = {item["key"] for item in FIELD_DEFINITIONS if item["required"]}

    for definition in FIELD_DEFINITIONS:
        key = definition["key"]
        source_header = mapping.get(key, "")
        if not isinstance(source_header, str):
            source_header = ""
        source_header = source_header.strip()
        supplied_header = source_header
        if source_header and source_header not in header_set:
            errors.append(_error("mapping_unknown_column", f"{definition['label']}映射的原始列“{source_header}”不存在", field=key))
            source_header = ""
        clean_mapping[key] = source_header
        if key in required_keys and not source_header and not supplied_header:
            errors.append(_error("mapping_required", f"必须映射{definition['label']}", field=key))

    used_headers = [header for header in clean_mapping.values() if header]
    for header in sorted({header for header in used_headers if used_headers.count(header) > 1}):
        errors.append(_error("mapping_duplicate_source", f"源列“{header}”不能映射到多个统一字段"))

    tasks: list[dict[str, Any]] = []
    project_pairs: list[tuple[str, str, int]] = []
    task_rows: dict[str, list[int]] = {}
    for raw in parsed.rows:
        row_number = raw["source_row"]
        values = raw["values"]
        get = lambda key: values.get(clean_mapping.get(key, "")) if clean_mapping.get(key) else None

        required_text = {}
        for key in ("project_id", "project_name", "task_id", "task_name", "owner"):
            value = _text(get(key))
            required_text[key] = value
            if not value:
                label = next(item["label"] for item in FIELD_DEFINITIONS if item["key"] == key)
                errors.append(_error("required_value_missing", f"第 {row_number} 行缺少{label}", row_number, key))

        start_date = _parse_date(get("start_date"))
        due_date = _parse_date(get("due_date"))
        if not start_date:
            errors.append(_error("invalid_date", f"第 {row_number} 行开始日期无效", row_number, "start_date"))
        if not due_date:
            errors.append(_error("invalid_date", f"第 {row_number} 行截止日期无效", row_number, "due_date"))
        if start_date and due_date and due_date < start_date:
            errors.append(_error("date_order", f"第 {row_number} 行截止日期早于开始日期", row_number, "due_date"))

        status_raw = _header_key(_text(get("status")))
        status = STATUS_ALIASES.get(status_raw)
        if not status:
            errors.append(_error("invalid_status", f"第 {row_number} 行状态不受支持", row_number, "status"))

        progress_value = get("progress_percent")
        progress = _parse_number(progress_value)
        if progress is not None and 0 <= progress <= 1:
            progress *= 100
        if progress_value not in (None, "") and (progress is None or not 0 <= progress <= 100):
            errors.append(_error("invalid_progress", f"第 {row_number} 行完成百分比必须在 0 到 100 之间", row_number, "progress_percent"))
            progress = None

        effort_value = get("effort_hours")
        effort = _parse_number(effort_value)
        if effort_value not in (None, "") and (effort is None or effort < 0):
            errors.append(_error("invalid_effort", f"第 {row_number} 行工作量必须是非负数", row_number, "effort_hours"))
            effort = None

        updated_value = get("updated_at")
        updated_at = _parse_datetime(updated_value)
        if updated_value not in (None, "") and not updated_at:
            errors.append(_error("invalid_updated_at", f"第 {row_number} 行更新时间无效", row_number, "updated_at"))

        task_id = required_text["task_id"]
        if task_id:
            task_rows.setdefault(task_id, []).append(row_number)
        project_pairs.append((required_text["project_id"], required_text["project_name"], row_number))
        tasks.append(
            {
                "source_row": row_number,
                **required_text,
                "start_date": start_date,
                "due_date": due_date,
                "status": status,
                "priority": _text(get("priority")) or None,
                "dependency_ids": _split_dependencies(get("dependency_ids")),
                "progress_percent": round(progress, 2) if progress is not None else None,
                "effort_hours": round(effort, 2) if effort is not None else None,
                "update_text": _text(get("update_text")) or None,
                "updated_at": updated_at,
            }
        )

    consistent_projects = {(project_id, project_name) for project_id, project_name, _ in project_pairs if project_id and project_name}
    if len(consistent_projects) > 1:
        errors.append(_error("project_inconsistent", "同一文件中的项目编号和项目名称必须一致"))

    for task_id, rows in task_rows.items():
        if len(rows) > 1:
            for duplicate_row in rows[1:]:
                errors.append(_error("duplicate_task_id", f"第 {duplicate_row} 行任务编号“{task_id}”重复", duplicate_row, "task_id"))

    known_task_ids = set(task_rows)
    dependencies: list[dict[str, Any]] = []
    updates: list[dict[str, Any]] = []
    for task in tasks:
        for dependency_id in task["dependency_ids"]:
            if dependency_id == task["task_id"]:
                errors.append(_error("self_dependency", f"第 {task['source_row']} 行任务不能依赖自己", task["source_row"], "dependency_ids"))
            elif dependency_id not in known_task_ids:
                errors.append(_error("dependency_not_found", f"第 {task['source_row']} 行前置任务“{dependency_id}”不存在", task["source_row"], "dependency_ids"))
            else:
                dependencies.append({"task_id": task["task_id"], "depends_on_task_id": dependency_id, "source_row": task["source_row"]})
        if task["update_text"] or task["updated_at"]:
            updates.append(
                {
                    "task_id": task["task_id"],
                    "source_row": task["source_row"],
                    "update_text": task["update_text"],
                    "updated_at": task["updated_at"],
                }
            )

    project_id = tasks[0]["project_id"] if tasks else ""
    project_name = tasks[0]["project_name"] if tasks else ""
    return {
        "valid": not errors,
        "mapping": clean_mapping,
        "errors": errors,
        "project": {"project_id": project_id, "project_name": project_name},
        "tasks": tasks,
        "dependencies": dependencies,
        "updates": updates,
        "summary": {
            "task_count": len(tasks),
            "dependency_count": len(dependencies),
            "update_count": len(updates),
            "error_count": len(errors),
        },
    }


def preview_payload(parsed: ParsedWorkbook, mapping: dict[str, Any], preview_id: str | None = None) -> dict[str, Any]:
    result = validate_and_normalize(parsed, mapping)
    return {
        "preview_id": preview_id,
        "source": {
            "source_name": parsed.source_name,
            "source_type": "xlsx",
            "sheet_name": parsed.sheet_name,
            "sha256": parsed.sha256,
        },
        "headers": parsed.headers,
        "field_definitions": FIELD_DEFINITIONS,
        **result,
        "tasks": result["tasks"][:MAX_PREVIEW_ROWS],
        "preview_truncated": len(result["tasks"]) > MAX_PREVIEW_ROWS,
    }


def save_pending_upload(content: bytes, source_name: str, pending_dir: Path) -> tuple[str, Path]:
    if not source_name.lower().endswith(".xlsx"):
        raise IngestionError("unsupported_file_type", "当前只支持 .xlsx 文件")
    if not content:
        raise IngestionError("empty_file", "XLSX 文件不能为空")
    if len(content) > MAX_XLSX_BYTES:
        raise IngestionError("file_too_large", "XLSX 文件不能超过 2 MB", 413)
    preview_id = uuid4().hex
    pending_dir.mkdir(parents=True, exist_ok=True)
    workbook_path = pending_dir / f"{preview_id}.xlsx"
    metadata_path = pending_dir / f"{preview_id}.json"
    workbook_path.write_bytes(content)
    metadata_path.write_text(json.dumps({"source_name": _safe_filename(source_name)}, ensure_ascii=False), encoding="utf-8")
    return preview_id, workbook_path


def load_pending(preview_id: str, pending_dir: Path) -> tuple[Path, str]:
    if not re.fullmatch(r"[0-9a-f]{32}", str(preview_id or "")):
        raise IngestionError("preview_not_found", "导入预览不存在或已经失效", 404)
    workbook_path = pending_dir / f"{preview_id}.xlsx"
    metadata_path = pending_dir / f"{preview_id}.json"
    try:
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError) as error:
        raise IngestionError("preview_not_found", "导入预览不存在或已经失效", 404) from error
    if not workbook_path.exists():
        raise IngestionError("preview_not_found", "导入预览不存在或已经失效", 404)
    return workbook_path, _safe_filename(metadata.get("source_name", ""))


def confirm_import(
    preview_id: str,
    mapping: dict[str, Any],
    pending_dir: Path,
    raw_dir: Path,
    normalized_dir: Path,
) -> dict[str, Any]:
    pending_path, source_name = load_pending(preview_id, pending_dir)
    parsed = parse_xlsx(pending_path, source_name)
    result = validate_and_normalize(parsed, mapping)
    if result["errors"]:
        raise IngestionError("validation_failed", "文件仍有校验错误，不能确认导入", 400)

    source_id = f"source-{parsed.sha256[:12]}"
    raw_dir.mkdir(parents=True, exist_ok=True)
    normalized_dir.mkdir(parents=True, exist_ok=True)
    raw_path = raw_dir / f"{source_id}-{_safe_filename(source_name)}"
    if raw_path.exists():
        raise IngestionError("duplicate_file", "相同内容的文件已经导入，请不要重复确认", 409)

    imported_at = datetime.now().astimezone().isoformat(timespec="seconds")
    import_id = f"import-{datetime.now().strftime('%Y%m%d%H%M%S')}-{parsed.sha256[:8]}"
    document = {
        "schema_version": SCHEMA_VERSION,
        "import_id": import_id,
        "source": {
            "source_id": source_id,
            "source_type": "xlsx",
            "source_name": source_name,
            "source_version": "v1",
            "sha256": parsed.sha256,
            "sheet_name": parsed.sheet_name,
            "imported_at": imported_at,
        },
        "mapping": result["mapping"],
        "project": result["project"],
        "tasks": result["tasks"],
        "dependencies": result["dependencies"],
        "updates": result["updates"],
        "summary": result["summary"],
    }
    normalized_path = normalized_dir / f"{import_id}.json"
    temporary_raw = raw_path.with_suffix(raw_path.suffix + ".tmp")
    temporary_json = normalized_path.with_suffix(".json.tmp")
    try:
        shutil.copyfile(pending_path, temporary_raw)
        temporary_json.write_text(json.dumps(document, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        temporary_raw.replace(raw_path)
        temporary_json.replace(normalized_path)
    finally:
        temporary_raw.unlink(missing_ok=True)
        temporary_json.unlink(missing_ok=True)

    pending_path.unlink(missing_ok=True)
    (pending_dir / f"{preview_id}.json").unlink(missing_ok=True)
    return {
        "message": "项目任务已确认导入",
        "import_id": import_id,
        "source": document["source"],
        "project": document["project"],
        "summary": document["summary"],
        "saved": {
            "raw": str(raw_path.relative_to(raw_dir.parent)).replace("\\", "/"),
            "normalized": str(normalized_path.relative_to(normalized_dir.parent)).replace("\\", "/"),
        },
    }
